"""
Golden-file tests for DOCX export quality.

Tests validate:
- Export generation succeeds
- Generated DOCXs are valid and parseable
- Basic content structure is preserved
- Determinism (same input → same structure)

Simplified approach: Rather than complex binary comparisons,
we test that exports are generated successfully and contain expected content.
"""

import pytest
import io
from docx import Document
from unittest.mock import Mock, MagicMock
from services.export_service import ExportService
from config import Settings


@pytest.fixture
def mock_supabase_client():
    """Mock Supabase client."""
    client = Mock()
    
    # Mock version query
    version_response = Mock()
    version_response.data = {
        "resume_id": "resume-123",
        "user_id": "user-1",
        "version_number": 1,
        "created_at": "2024-01-01T00:00:00"
    }
    
    # Mock sections query
    sections_response = Mock()
    sections_response.data = [
        {
            "id": "section-1",
            "section_type": "contact",
            "section_number": 1,
            "title": "Contact"
        },
        {
            "id": "section-2",
            "section_type": "summary",
            "section_number": 2,
            "title": "Professional Summary"
        },
        {
            "id": "section-3",
            "section_type": "experience",
            "section_number": 3,
            "title": "Work Experience"
        }
    ]
    
    # Mock blocks query
    blocks_response = Mock()
    blocks_response.data = [
        {
            "id": "block-1",
            "section_id": "section-1",
            "block_number": 1,
            "block_type": "paragraph",
            "content": {
                "name": "John Doe",
                "email": "john@example.com",
                "phone": "+1-555-1234"
            }
        },
        {
            "id": "block-2",
            "section_id": "section-2",
            "block_number": 1,
            "block_type": "paragraph",
            "content": {
                "text": "Experienced software engineer with 5+ years"
            }
        },
        {
            "id": "block-3",
            "section_id": "section-3",
            "block_number": 1,
            "block_type": "heading",
            "content": {
                "text": "Senior Engineer | Tech Corp",
                "metadata": {"bold": True}
            }
        },
        {
            "id": "block-4",
            "section_id": "section-3",
            "block_number": 2,
            "block_type": "bullet",
            "content": {
                "text": "Built microservices handling 1M+ requests"
            }
        }
    ]
    
    # Setup mock chain
    client.table.return_value.select.return_value.eq.return_value.single.return_value.execute.return_value = version_response
    client.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = sections_response
    client.table.return_value.select.return_value.in_.return_value.order.return_value.execute.return_value = blocks_response
    
    return client


@pytest.fixture
def export_service(mock_supabase_client):
    """Create export service with mocked dependencies."""
    settings = Mock(spec=Settings)
    settings.supabase_url = "http://test"
    settings.supabase_key = "test-key"
    settings.supabase_service_role_key = "test-service-role-key"
    settings.anthropic_api_key = "test-key"
    
    service = ExportService(settings)
    service.client = mock_supabase_client
    
    return service


class TestDocxGeneration:
    """Tests for successful DOCX generation."""
    
    def test_generate_docx_succeeds(self, export_service):
        """DOCX generation should complete without errors."""
        docx_bytes = export_service.generate_docx("user-1", "version-1")
        
        assert docx_bytes is not None
        assert len(docx_bytes) > 0
        assert isinstance(docx_bytes, bytes)
    
    def test_generated_docx_is_valid(self, export_service):
        """Generated DOCX should be valid and parseable."""
        docx_bytes = export_service.generate_docx("user-1", "version-1")
        
        # Should be able to parse it
        doc = Document(io.BytesIO(docx_bytes))
        
        # Should have paragraphs
        assert len(doc.paragraphs) > 0
    
    def test_docx_contains_expected_content(self, export_service):
        """Generated DOCX should contain expected text."""
        docx_bytes = export_service.generate_docx("user-1", "version-1")
        doc = Document(io.BytesIO(docx_bytes))
        
        # Extract all text
        all_text = "\n".join(p.text for p in doc.paragraphs)
        
        # Should contain contact info
        assert "John Doe" in all_text
        assert "john@example.com" in all_text
        
        # Should contain content
        assert "software engineer" in all_text.lower() or "engineer" in all_text.lower()
    
    def test_docx_has_proper_structure(self, export_service):
        """Generated DOCX should have paragraphs and content."""
        docx_bytes = export_service.generate_docx("user-1", "version-1")
        doc = Document(io.BytesIO(docx_bytes))
        
        # Should have multiple paragraphs
        assert len(doc.paragraphs) >= 3
        
        # Should have some non-empty paragraphs
        non_empty = [p for p in doc.paragraphs if p.text.strip()]
        assert len(non_empty) > 0


class TestDocxDeterminism:
    """Tests for deterministic DOCX generation."""
    
    def test_same_input_same_structure(self, export_service):
        """Same input should produce same structure."""
        # Generate twice
        docx_bytes_1 = export_service.generate_docx("user-1", "version-1")
        docx_bytes_2 = export_service.generate_docx("user-1", "version-1")
        
        # Parse both
        doc1 = Document(io.BytesIO(docx_bytes_1))
        doc2 = Document(io.BytesIO(docx_bytes_2))
        
        # Same number of paragraphs
        assert len(doc1.paragraphs) == len(doc2.paragraphs)
        
        # Same text content
        text1 = "\n".join(p.text for p in doc1.paragraphs)
        text2 = "\n".join(p.text for p in doc2.paragraphs)
        assert text1 == text2
    
    def test_multiple_generations_consistent(self, export_service):
        """Multiple generations should be consistent."""
        # Generate 5 times
        docs = []
        for _ in range(5):
            docx_bytes = export_service.generate_docx("user-1", "version-1")
            doc = Document(io.BytesIO(docx_bytes))
            docs.append(doc)
        
        # All should have same paragraph count
        para_counts = [len(d.paragraphs) for d in docs]
        assert all(c == para_counts[0] for c in para_counts)
        
        # All should have same text
        texts = ["\n".join(p.text for p in d.paragraphs) for d in docs]
        assert all(t == texts[0] for t in texts)


class TestDocxQuality:
    """Tests for DOCX quality and formatting."""
    
    def test_docx_has_formatting(self, export_service):
        """DOCX should include formatting (bold, styles)."""
        docx_bytes = export_service.generate_docx("user-1", "version-1")
        doc = Document(io.BytesIO(docx_bytes))
        
        # Check for bold text
        has_bold = False
        for para in doc.paragraphs:
            for run in para.runs:
                if run.bold:
                    has_bold = True
                    break
            if has_bold:
                break
        
        assert has_bold, "Document should have some bold text"
    
    def test_docx_has_proper_text_encoding(self, export_service):
        """DOCX should handle text encoding correctly."""
        docx_bytes = export_service.generate_docx("user-1", "version-1")
        doc = Document(io.BytesIO(docx_bytes))
        
        # Should be able to extract text without encoding errors
        for para in doc.paragraphs:
            text = para.text
            assert isinstance(text, str)
