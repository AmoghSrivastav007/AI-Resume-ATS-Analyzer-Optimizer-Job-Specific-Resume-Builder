"""
Tests for Step 3: Structural Extraction (Rule-based parsing).

Tests deterministic PDF/DOCX parsing functions.
These are pure functions with zero variance - same input always produces same output.
"""

import io
import pytest
from docx import Document
from docx.shared import Pt
import fitz  # PyMuPDF

from services.parser.structural import (
    extract_structure,
    _extract_pdf_structure,
    _extract_docx_structure,
    StructuralBlock,
    StructuralDocument,
)
from services.file_validation import PDF_MIME, DOCX_MIME


class TestStructuralBlock:
    """Tests for StructuralBlock dataclass."""
    
    def test_structural_block_immutability(self):
        """StructuralBlock should be immutable (frozen dataclass)."""
        block = StructuralBlock(
            text="Test text",
            block_type="line",
            page=1,
            font_size=12.0,
            sort_order=0
        )
        
        # Should not be able to modify frozen dataclass
        with pytest.raises(AttributeError):
            block.text = "Modified"
    
    def test_structural_block_defaults(self):
        """StructuralBlock should have proper defaults."""
        block = StructuralBlock(text="Test", block_type="line")
        
        assert block.text == "Test"
        assert block.block_type == "line"
        assert block.page is None
        assert block.bbox is None
        assert block.font_name is None
        assert block.font_size is None
        assert block.sort_order == 0


class TestDocxExtraction:
    """Tests for DOCX structural extraction."""
    
    def test_extract_docx_structure_basic(self):
        """Extract basic paragraphs from DOCX."""
        # Create minimal DOCX in memory
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        doc.add_paragraph("Third paragraph")
        
        # Save to bytes
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        
        # Extract
        result = _extract_docx_structure(content)
        
        # Verify
        assert result.mime_type == DOCX_MIME
        assert len(result.blocks) == 3
        assert result.blocks[0].text == "First paragraph"
        assert result.blocks[1].text == "Second paragraph"
        assert result.blocks[2].text == "Third paragraph"
        assert all(b.block_type == "paragraph" for b in result.blocks)
        
        # Verify sort order
        assert result.blocks[0].sort_order == 0
        assert result.blocks[1].sort_order == 1
        assert result.blocks[2].sort_order == 2
        
        # Determinism check - run again
        result2 = _extract_docx_structure(content)
        assert len(result2.blocks) == len(result.blocks)
        assert all(b1.text == b2.text for b1, b2 in zip(result.blocks, result2.blocks))
    
    def test_extract_docx_structure_headings(self):
        """Detect headings vs paragraphs in DOCX."""
        doc = Document()
        doc.add_heading("Main Heading", level=1)
        doc.add_paragraph("Normal paragraph")
        doc.add_heading("Subheading", level=2)
        
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        
        result = _extract_docx_structure(content)
        
        assert len(result.blocks) == 3
        assert result.blocks[0].block_type == "heading"  # Heading 1
        assert result.blocks[1].block_type == "paragraph"  # Normal
        assert result.blocks[2].block_type == "heading"  # Heading 2
        
        # Determinism
        result2 = _extract_docx_structure(content)
        assert [b.block_type for b in result.blocks] == [b.block_type for b in result2.blocks]
    
    def test_extract_docx_structure_empty_paragraphs(self):
        """Empty paragraphs should be skipped."""
        doc = Document()
        doc.add_paragraph("First")
        doc.add_paragraph("")  # Empty
        doc.add_paragraph("   ")  # Whitespace only
        doc.add_paragraph("Second")
        
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        
        result = _extract_docx_structure(content)
        
        # Should only have 2 blocks (empty ones skipped)
        assert len(result.blocks) == 2
        assert result.blocks[0].text == "First"
        assert result.blocks[1].text == "Second"
        
        # Determinism
        result2 = _extract_docx_structure(content)
        assert len(result2.blocks) == 2
    
    def test_extract_docx_structure_font_info(self):
        """Extract font information from DOCX."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Text with font")
        run.font.name = "Arial"
        run.font.size = Pt(14)
        
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        
        result = _extract_docx_structure(content)
        
        assert len(result.blocks) == 1
        block = result.blocks[0]
        assert block.font_name == "Arial"
        assert block.font_size == 14.0
        
        # Determinism
        result2 = _extract_docx_structure(content)
        assert result2.blocks[0].font_name == block.font_name
        assert result2.blocks[0].font_size == block.font_size
    
    def test_extract_docx_plain_text(self):
        """Verify plain text extraction."""
        doc = Document()
        doc.add_paragraph("Line 1")
        doc.add_paragraph("Line 2")
        doc.add_paragraph("Line 3")
        
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        
        result = _extract_docx_structure(content)
        
        # Plain text should contain all lines
        assert "Line 1" in result.plain_text
        assert "Line 2" in result.plain_text
        assert "Line 3" in result.plain_text
        
        # Determinism
        result2 = _extract_docx_structure(content)
        assert result2.plain_text == result.plain_text


class TestPdfExtraction:
    """Tests for PDF structural extraction."""
    
    def test_extract_pdf_structure_basic(self):
        """Extract basic text blocks from PDF."""
        # Create minimal PDF in memory using PyMuPDF
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "First line of text")
        page.insert_text((72, 100), "Second line of text")
        
        # Save to bytes
        content = doc.tobytes()
        doc.close()
        
        # Extract
        result = _extract_pdf_structure(content)
        
        # Verify
        assert result.mime_type == PDF_MIME
        assert len(result.blocks) >= 2  # At least 2 text blocks
        
        # Find our text in blocks
        texts = [b.text for b in result.blocks]
        assert any("First line" in t for t in texts)
        assert any("Second line" in t for t in texts)
        
        # All blocks should have page number
        assert all(b.page == 1 for b in result.blocks)
        
        # Determinism check
        result2 = _extract_pdf_structure(content)
        assert len(result2.blocks) == len(result.blocks)
    
    def test_extract_pdf_structure_multi_page(self):
        """Extract from multi-page PDF."""
        doc = fitz.open()
        
        # Page 1
        page1 = doc.new_page()
        page1.insert_text((72, 72), "Page one text")
        
        # Page 2
        page2 = doc.new_page()
        page2.insert_text((72, 72), "Page two text")
        
        content = doc.tobytes()
        doc.close()
        
        result = _extract_pdf_structure(content)
        
        # Should have blocks from both pages
        page_numbers = {b.page for b in result.blocks}
        assert 1 in page_numbers
        assert 2 in page_numbers
        
        # Verify text from each page
        texts = [b.text for b in result.blocks]
        assert any("Page one" in t for t in texts)
        assert any("Page two" in t for t in texts)
        
        # Determinism
        result2 = _extract_pdf_structure(content)
        assert {b.page for b in result2.blocks} == page_numbers
    
    def test_extract_pdf_structure_font_detection(self):
        """Font name and size should be extracted."""
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Text with font", fontsize=16)
        
        content = doc.tobytes()
        doc.close()
        
        result = _extract_pdf_structure(content)
        
        assert len(result.blocks) > 0
        # At least one block should have font info
        has_font = any(b.font_size is not None for b in result.blocks)
        assert has_font
        
        # Determinism
        result2 = _extract_pdf_structure(content)
        assert len(result2.blocks) == len(result.blocks)
    
    def test_extract_pdf_structure_empty_lines(self):
        """Empty lines should be skipped."""
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Text")
        page.insert_text((72, 100), "")  # Empty
        page.insert_text((72, 128), "More text")
        
        content = doc.tobytes()
        doc.close()
        
        result = _extract_pdf_structure(content)
        
        # Empty lines should be skipped
        texts = [b.text for b in result.blocks]
        assert all(t.strip() != "" for t in texts)
        
        # Determinism
        result2 = _extract_pdf_structure(content)
        assert len(result2.blocks) == len(result.blocks)
    
    def test_extract_pdf_plain_text(self):
        """Verify plain text extraction."""
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Line 1")
        page.insert_text((72, 100), "Line 2")
        
        content = doc.tobytes()
        doc.close()
        
        result = _extract_pdf_structure(content)
        
        # Plain text should contain lines
        assert "Line 1" in result.plain_text
        assert "Line 2" in result.plain_text
        
        # Determinism
        result2 = _extract_pdf_structure(content)
        assert result2.plain_text == result.plain_text


class TestExtractStructureDispatcher:
    """Tests for extract_structure() dispatcher function."""
    
    def test_extract_structure_pdf(self):
        """extract_structure should route PDF correctly."""
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "PDF text")
        content = doc.tobytes()
        doc.close()
        
        result = extract_structure(content, PDF_MIME)
        
        assert result.mime_type == PDF_MIME
        assert len(result.blocks) > 0
        
        # Determinism
        result2 = extract_structure(content, PDF_MIME)
        assert result2.mime_type == result.mime_type
        assert len(result2.blocks) == len(result.blocks)
    
    def test_extract_structure_docx(self):
        """extract_structure should route DOCX correctly."""
        doc = Document()
        doc.add_paragraph("DOCX text")
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        
        result = extract_structure(content, DOCX_MIME)
        
        assert result.mime_type == DOCX_MIME
        assert len(result.blocks) == 1
        assert result.blocks[0].text == "DOCX text"
        
        # Determinism
        result2 = extract_structure(content, DOCX_MIME)
        assert result2.mime_type == result.mime_type
        assert result2.blocks[0].text == result.blocks[0].text
    
    def test_extract_structure_unsupported_type(self):
        """extract_structure should raise error for unsupported types."""
        content = b"some content"
        
        with pytest.raises(ValueError, match="Unsupported mime type"):
            extract_structure(content, "text/plain")
    
    def test_extract_structure_determinism(self):
        """Verify determinism: same input -> same output."""
        # Create DOCX
        doc = Document()
        doc.add_paragraph("Determinism test")
        doc.add_paragraph("Second line")
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        
        # Extract multiple times
        result1 = extract_structure(content, DOCX_MIME)
        result2 = extract_structure(content, DOCX_MIME)
        result3 = extract_structure(content, DOCX_MIME)
        
        # All should be identical
        assert len(result1.blocks) == len(result2.blocks) == len(result3.blocks)
        assert result1.plain_text == result2.plain_text == result3.plain_text
        
        for b1, b2, b3 in zip(result1.blocks, result2.blocks, result3.blocks):
            assert b1.text == b2.text == b3.text
            assert b1.block_type == b2.block_type == b3.block_type
            assert b1.sort_order == b2.sort_order == b3.sort_order
