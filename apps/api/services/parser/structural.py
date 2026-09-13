import io
from dataclasses import dataclass, field

import fitz
import mammoth
from docx import Document

from services.file_validation import DOCX_MIME, PDF_MIME


@dataclass(frozen=True)
class StructuralBlock:
    text: str
    block_type: str
    page: int | None = None
    bbox: list[float] | None = None
    font_name: str | None = None
    font_size: float | None = None
    sort_order: int = 0


@dataclass
class StructuralDocument:
    mime_type: str
    plain_text: str
    blocks: list[StructuralBlock] = field(default_factory=list)


def extract_structure(content: bytes, mime_type: str) -> StructuralDocument:
    if mime_type == PDF_MIME:
        return _extract_pdf_structure(content)
    if mime_type == DOCX_MIME:
        return _extract_docx_structure(content)
    raise ValueError(f"Unsupported mime type for structural extraction: {mime_type}")


def _extract_pdf_structure(content: bytes) -> StructuralDocument:
    document = fitz.open(stream=content, filetype="pdf")
    blocks: list[StructuralBlock] = []
    plain_parts: list[str] = []
    sort_order = 0

    try:
        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            page_dict = page.get_text("dict")
            for block in page_dict.get("blocks", []):
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    line_text = "".join(span.get("text", "") for span in spans).strip()
                    if not line_text:
                        continue
                    plain_parts.append(line_text)
                    primary_span = spans[0] if spans else {}
                    blocks.append(
                        StructuralBlock(
                            text=line_text,
                            block_type="line",
                            page=page_index + 1,
                            bbox=list(block.get("bbox", [])) or None,
                            font_name=primary_span.get("font"),
                            font_size=float(primary_span.get("size", 0)) or None,
                            sort_order=sort_order,
                        )
                    )
                    sort_order += 1
    finally:
        document.close()

    plain_text = "\n".join(plain_parts)
    return StructuralDocument(mime_type=PDF_MIME, plain_text=plain_text, blocks=blocks)


def _extract_docx_structure(content: bytes) -> StructuralDocument:
    document = Document(io.BytesIO(content))
    blocks: list[StructuralBlock] = []
    plain_parts: list[str] = []
    sort_order = 0

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else "Normal"
        block_type = "heading" if "Heading" in style_name else "paragraph"
        run_fonts = [
            run.font.name
            for run in paragraph.runs
            if run.font is not None and run.font.name
        ]
        run_sizes = [
            float(run.font.size.pt)
            for run in paragraph.runs
            if run.font is not None and run.font.size is not None
        ]
        blocks.append(
            StructuralBlock(
                text=text,
                block_type=block_type,
                font_name=run_fonts[0] if run_fonts else None,
                font_size=run_sizes[0] if run_sizes else None,
                sort_order=sort_order,
            )
        )
        plain_parts.append(text)
        sort_order += 1

    mammoth_result = mammoth.extract_raw_text(io.BytesIO(content))
    mammoth_text = mammoth_result.value.strip()
    plain_text = mammoth_text or "\n".join(plain_parts)

    return StructuralDocument(mime_type=DOCX_MIME, plain_text=plain_text, blocks=blocks)
